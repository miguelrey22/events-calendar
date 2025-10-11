"""
Events Calendar AKS - Al Kamel Management
Sistema Completo de Gestión Visual de Eventos

Versión: 3.1 - Con detección de conexiones de viaje
Autor: Claude AI para Alkamel Management
Fecha: 11/10/2025
"""

import os
import json
import pandas as pd
import numpy as np
from datetime import datetime, timedelta, date
import requests
import msal
import threading
import schedule
import time
from flask import Flask, render_template, request, jsonify
import logging
from typing import Dict, List, Optional, Tuple
import sqlite3
from io import BytesIO
import warnings
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from collections import defaultdict

warnings.filterwarnings('ignore')

# Configuración de logging
logging.basicConfig(
    level=logging.INFO,
    format='%(asctime)s - %(levelname)s - %(message)s'
)
logger = logging.getLogger(__name__)

class EventsCalendarAKS:
    """Sistema de calendario visual para Al Kamel Management"""
    
    def __init__(self, config: Dict):
        self.config = config
        self.airtable_token = config.get('airtable_token')
        self.airtable_base_id = config.get('airtable_base_id', 'app4p2TY96NofXW4u')
        
        # Microsoft Graph
        self.tenant_id = config.get('tenant_id')
        self.client_id = config.get('client_id')
        self.client_secret = config.get('client_secret')
        self.sharepoint_site_url = config.get('sharepoint_site_url')
        
        # Configuración
        self.auto_update_interval = config.get('auto_update_interval', 15)
        self.max_retries = 3
        self.timeout_seconds = 30
        
        # Cache
        self.cache = {}
        self.cache_expiry = {}
        
        # Headers Airtable
        self.airtable_headers = {
            'Authorization': f'Bearer {self.airtable_token}',
            'Content-Type': 'application/json'
        }
        
        # MSAL
        self.msal_app = self._init_msal()
        self.graph_token = None
        
        # Colores por SET
        self.color_mapping = {
            'SET 1': '#FF6B6B',
            'SET 2': '#4ECDC4',
            'SET 3': '#45B7D1',
            'MICROSET': '#96CEB4',
            'SET RW': '#FFEAA7',
            'EVENTS 3': '#DDA0DD',
            'EVENTS 4': '#98D8C8',
            'EVENTS 5': '#F7DC6F',
            'EVENTS 6': '#BB8FCE',
            'SCER': '#F8C471',
            'CERVH': '#85C1E9',
            'SET6': '#F1C40F',
            'SET 5': '#27AE60',
            'default': '#BDC3C7'
        }
        
        # Mapeo campeonatos a SETs
        self.championship_to_set = {
            'WEC': 'SET 1', 'FIA': 'SET 1',
            'CIRCUITCAT': 'SET 2', 'KATEYAMA': 'SET 2',
            'FERRARI': 'SET 3', 'MCLAREN': 'SET 3',
            'ELMS': 'SET 5',
            'F4': 'SET RW', 'E3': 'SET RW', 'GSERIES': 'SET RW',
            'E1': 'SET6',
            'SCER': 'SCER',
            'CERVH': 'CERVH'
        }
        
        logger.info("✅ Events Calendar AKS inicializado")
    
    def _init_msal(self):
        """Inicializar MSAL"""
        try:
            if not all([self.tenant_id, self.client_id, self.client_secret]):
                logger.warning("⚠️ Credenciales de Microsoft Graph incompletas")
                return None
            
            return msal.ConfidentialClientApplication(
                client_id=self.client_id,
                client_credential=self.client_secret,
                authority=f"https://login.microsoftonline.com/{self.tenant_id}"
            )
        except Exception as e:
            logger.error(f"❌ Error inicializando MSAL: {str(e)}")
            return None
    
    def get_graph_token(self) -> Optional[str]:
        """Obtener token Microsoft Graph"""
        if not self.msal_app:
            return None
        
        try:
            result = self.msal_app.acquire_token_for_client(
                scopes=["https://graph.microsoft.com/.default"]
            )
            
            if "access_token" in result:
                self.graph_token = result["access_token"]
                return self.graph_token
            else:
                logger.error(f"❌ Error token: {result.get('error_description')}")
                return None
        except Exception as e:
            logger.error(f"❌ Excepción obteniendo token: {str(e)}")
            return None
    
    def get_airtable_data(self, table_name: str) -> List[Dict]:
        """Obtener datos de Airtable con cache"""
        cache_key = f"airtable_{table_name}"
        
        if cache_key in self.cache:
            if datetime.now() < self.cache_expiry.get(cache_key, datetime.min):
                logger.info(f"📦 Usando cache para {table_name}")
                return self.cache[cache_key]
        
        table_ids = {
            'EVENTS': 'tblVb1BuNKkUoS96b',
            'EVENTS RESERVATIONS': 'tbllmzrlZvphVWaP7',
            'Employee directory': 'tblzwiTaABBdqaJ3G',
            'GUARDIAS': 'tblZtKR9x67vxayAF'
        }
        
        table_id = table_ids.get(table_name, table_name)
        url = f"https://api.airtable.com/v0/{self.airtable_base_id}/{table_id}"
        
        all_records = []
        
        try:
            params = {'pageSize': 100}
            
            while True:
                response = requests.get(
                    url, 
                    headers=self.airtable_headers, 
                    params=params,
                    timeout=self.timeout_seconds
                )
                
                if response.status_code == 200:
                    data = response.json()
                    all_records.extend(data.get('records', []))
                    
                    if 'offset' in data:
                        params['offset'] = data['offset']
                    else:
                        break
                else:
                    logger.error(f"❌ Error Airtable {table_name}: {response.status_code}")
                    break
            
            if all_records:
                self.cache[cache_key] = all_records
                self.cache_expiry[cache_key] = datetime.now() + timedelta(minutes=5)
                logger.info(f"📊 Obtenidos {len(all_records)} registros de {table_name}")
                
            return all_records
            
        except Exception as e:
            logger.error(f"❌ Excepción obteniendo datos de {table_name}: {str(e)}")
            return []
    
    def detect_conflicts(self, events: List[Dict]) -> Tuple[List[Dict], Dict]:
        """Detectar conflictos de personal"""
        conflicts = []
        employee_timelines = defaultdict(list)
        
        for event in events:
            for reservation in event['reservations']:
                employee_name = reservation['employee']
                employee_timelines[employee_name].append({
                    'event': event['event_name'],
                    'event_id': event['event_id'],
                    'from': reservation['from_date'],
                    'to': reservation['to_date'],
                    'city': event['city']
                })
        
        conflict_details = {}
        for employee, timeline in employee_timelines.items():
            timeline.sort(key=lambda x: x['from'])
            
            for i in range(len(timeline)):
                for j in range(i + 1, len(timeline)):
                    event1 = timeline[i]
                    event2 = timeline[j]
                    
                    if event1['to'] >= event2['from']:
                        conflict_key = f"{employee}_{event1['event_id']}_{event2['event_id']}"
                        if conflict_key not in conflict_details:
                            conflicts.append({
                                'employee': employee,
                                'event1': event1['event'],
                                'event2': event2['event'],
                                'city1': event1['city'],
                                'city2': event2['city'],
                                'dates': f"{event1['from'].strftime('%d/%m')} - {event2['to'].strftime('%d/%m')}"
                            })
                            conflict_details[conflict_key] = True
        
        logger.info(f"⚠️ Detectados {len(conflicts)} conflictos")
        return conflicts, employee_timelines
    
    def detect_travel_connections(self, events: List[Dict]) -> Dict:
        """
        Detectar qué personal viene de un evento la semana anterior o va a otro la semana siguiente
        """
        travel_connections = {}
        
        # Crear diccionario de eventos por empleado
        employee_events = defaultdict(list)
        
        for event in events:
            for reservation in event['reservations']:
                employee_name = reservation['employee']
                employee_events[employee_name].append({
                    'event_id': event['event_id'],
                    'event_name': event['event_name'],
                    'from_date': event['from_date'],
                    'to_date': event['to_date'],
                    'city': event['city']
                })
        
        # Para cada evento, detectar conexiones de viaje
        for event in events:
            event_connections = {
                'people_with_travel': [],
                'from_previous': [],
                'to_next': []
            }
            
            for reservation in event['reservations']:
                employee_name = reservation['employee']
                current_event_start = event['from_date']
                current_event_end = event['to_date']
                
                # Buscar eventos del mismo empleado
                emp_events = employee_events[employee_name]
                
                has_connection = False
                
                for other_event in emp_events:
                    if other_event['event_id'] == event['event_id']:
                        continue
                    
                    # Evento anterior (termina hasta 7 días antes)
                    days_between_prev = (current_event_start - other_event['to_date']).days
                    if 0 < days_between_prev <= 7:
                        event_connections['from_previous'].append({
                            'employee': employee_name,
                            'previous_event': other_event['event_name'],
                            'previous_city': other_event['city'],
                            'days_gap': days_between_prev
                        })
                        has_connection = True
                    
                    # Evento siguiente (empieza hasta 7 días después)
                    days_between_next = (other_event['from_date'] - current_event_end).days
                    if 0 < days_between_next <= 7:
                        event_connections['to_next'].append({
                            'employee': employee_name,
                            'next_event': other_event['event_name'],
                            'next_city': other_event['city'],
                            'days_gap': days_between_next
                        })
                        has_connection = True
                
                if has_connection:
                    event_connections['people_with_travel'].append(employee_name)
            
            travel_connections[event['event_id']] = event_connections
        
        logger.info(f"✈️ Detectadas conexiones de viaje para {len(travel_connections)} eventos")
        return travel_connections
    
    def find_available_staff(self, start_date: date, end_date: date, role_filter: str = None) -> List[Dict]:
        """Buscar personal disponible en un rango de fechas"""
        
        employees_data = self.get_airtable_data('Employee directory')
        reservations_data = self.get_airtable_data('EVENTS RESERVATIONS')
        
        # Emails genéricos a excluir (no son personas físicas)
        generic_emails = [
            'info@', 'admin@', 'operations@', 'contact@', 'support@',
            'hello@', 'office@', 'general@', 'staff@', 'team@'
        ]
        
        # Nombres falsos/placeholders a excluir
        fake_names = [
            'airtable.user1', 
            'tba', 
            'tbc',
            'to be announced',
            'to be confirmed',
            'por confirmar',
            'por anunciar',
            'pendiente'
        ]
        
        available_staff = []
        
        for emp_record in employees_data:
            emp_fields = emp_record.get('fields', {})
            emp_name = emp_fields.get('Name', 'Sin nombre')
            emp_email = emp_fields.get('EMAIL', '')
            emp_role = emp_fields.get('POSITION', '')
            
            # FILTRO 1: Excluir CUALQUIER nombre que contenga @ (es un email, no una persona)
            if '@' in emp_name:
                logger.debug(f"Excluido email como nombre: {emp_name}")
                continue
            
            # FILTRO 2: Excluir nombres falsos/placeholders
            emp_name_lower = emp_name.lower().strip()
            if any(fake_name in emp_name_lower for fake_name in fake_names):
                logger.debug(f"Excluido nombre falso/placeholder: {emp_name}")
                continue
            
            # FILTRO 3: Excluir si el nombre es demasiado corto (probablemente no es persona real)
            if len(emp_name.strip()) < 3:
                logger.debug(f"Excluido nombre muy corto: {emp_name}")
                continue
            
            # FILTRO 4: Excluir nombres que son claramente genéricos
            generic_names = ['operations', 'admin', 'info', 'contact', 'support', 'office', 'staff', 'team', 'general']
            if any(generic.lower() in emp_name.lower() for generic in generic_names):
                logger.debug(f"Excluido nombre genérico: {emp_name}")
                continue
            
            # FILTRO 5: Filtrar por rol si se especifica
            if role_filter and role_filter.lower() not in emp_role.lower():
                continue
            
            is_available = True
            last_event_date = None
            total_events = 0
            sets_experience = set()
            
            for res_record in reservations_data:
                res_fields = res_record.get('fields', {})
                
                emp_link = res_fields.get('Employee directory', [])
                if emp_record['id'] not in emp_link:
                    continue
                
                total_events += 1
                
                event_name = res_fields.get('Name (from EVENT)', [''])[0] if res_fields.get('Name (from EVENT)') else ''
                for key in self.championship_to_set.keys():
                    if key in event_name.upper():
                        sets_experience.add(self.championship_to_set[key])
                        break
                
                if 'FROM' in res_fields and 'TO' in res_fields:
                    try:
                        res_start = datetime.strptime(res_fields['FROM'], '%Y-%m-%d').date()
                        res_end = datetime.strptime(res_fields['TO'], '%Y-%m-%d').date()
                        
                        if not last_event_date or res_end > last_event_date:
                            last_event_date = res_end
                        
                        if not (res_end < start_date or res_start > end_date):
                            is_available = False
                            break
                    except:
                        continue
            
            if is_available:
                days_available = 365 - (total_events * 3)
                
                available_staff.append({
                    'name': emp_name,
                    'email': emp_email,
                    'role': emp_role,
                    'total_events': total_events,
                    'sets_experience': list(sets_experience),
                    'last_event': last_event_date.strftime('%d/%m/%Y') if last_event_date else 'Nunca',
                    'days_available': max(0, days_available)
                })
        
        available_staff.sort(key=lambda x: x['total_events'], reverse=True)
        
        logger.info(f"✅ Encontrados {len(available_staff)} empleados disponibles (filtrados nombres falsos)")
        return available_staff
    
    def process_motorsport_data(self) -> Dict:
        """Procesar datos completos"""
        logger.info("🔄 Procesando datos...")
        
        events_data = self.get_airtable_data('EVENTS')
        reservations_data = self.get_airtable_data('EVENTS RESERVATIONS')
        
        if not events_data:
            logger.error("❌ No se encontraron eventos")
            return {}
        
        start_date = datetime.now().date()
        end_date = start_date + timedelta(days=365)
        
        processed_events = []
        unassigned_events = []
        stats = {
            'total_events': 0,
            'confirmed_events': 0,
            'unassigned_events': 0,
            'total_reservations': 0,
            'remote_assignments': 0,
            'events_by_set': {},
            'events_by_month': defaultdict(int),
            'events_by_week': defaultdict(list),
            'critical_dates': []
        }
        
        for event_record in events_data:
            fields = event_record.get('fields', {})
            
            if 'From' not in fields or 'To' not in fields:
                continue
            
            try:
                event_start = datetime.strptime(fields['From'], '%Y-%m-%d').date()
                event_end = datetime.strptime(fields['To'], '%Y-%m-%d').date()
            except:
                continue
            
            if event_start > end_date or event_end < start_date:
                continue
            
            championship = fields.get('CAMPEONATO-CIRCUITO-ENTIDAD (from CHAMPIONSHIP)', [''])[0] if fields.get('CAMPEONATO-CIRCUITO-ENTIDAD (from CHAMPIONSHIP)') else ''
            set_name = self._determine_set(championship)
            
            confirmed = fields.get('CONFIRMED', False)
            coordinator = fields.get('Name (from Event Coordinator)', [''])[0] if fields.get('Name (from Event Coordinator)') else 'Sin coordinador'
            
            event_reservations = []
            for res_record in reservations_data:
                res_fields = res_record.get('fields', {})
                if res_fields.get('EVENT', [''])[0] == event_record['id']:
                    if 'FROM' in res_fields and 'TO' in res_fields:
                        try:
                            res_start = datetime.strptime(res_fields['FROM'], '%Y-%m-%d').date()
                            res_end = datetime.strptime(res_fields['TO'], '%Y-%m-%d').date()
                            
                            employee_name = res_fields.get('Name (from Employee directory)', ['Sin asignar'])[0] if res_fields.get('Name (from Employee directory)') else 'Sin asignar'
                            is_remote = res_fields.get('REMOTE', False)
                            
                            event_reservations.append({
                                'employee': employee_name,
                                'from_date': res_start,
                                'to_date': res_end,
                                'remote': is_remote,
                                'days': (res_end - res_start).days + 1
                            })
                            
                            if is_remote:
                                stats['remote_assignments'] += 1
                        except:
                            continue
            
            week_num = event_start.isocalendar()[1]
            month_key = event_start.strftime('%Y-%m')
            
            event_entry = {
                'event_id': event_record['id'],
                'event_name': fields.get('EVENT NAME', 'Sin nombre'),
                'city': fields.get('EVENT CITY', ''),
                'championship': championship,
                'set_name': set_name,
                'color': self.color_mapping.get(set_name, self.color_mapping['default']),
                'coordinator': coordinator,
                'confirmed': confirmed,
                'from_date': event_start,
                'to_date': event_end,
                'duration_days': (event_end - event_start).days + 1,
                'reservations': event_reservations,
                'employees_count': len(event_reservations),
                'needs_attention': len(event_reservations) == 0 and confirmed,
                'week_number': week_num,
                'month': month_key
            }
            
            processed_events.append(event_entry)
            
            stats['total_events'] += 1
            stats['total_reservations'] += len(event_reservations)
            stats['events_by_month'][month_key] += 1
            stats['events_by_week'][week_num].append(event_entry)
            
            if confirmed:
                stats['confirmed_events'] += 1
            
            if len(event_reservations) == 0 and confirmed:
                stats['unassigned_events'] += 1
                unassigned_events.append(event_entry)
            
            stats['events_by_set'][set_name] = stats['events_by_set'].get(set_name, 0) + 1
            
            if (event_start - start_date).days <= 7 and len(event_reservations) == 0 and confirmed:
                stats['critical_dates'].append(event_entry)
        
        processed_events.sort(key=lambda x: x['from_date'])
        
        conflicts, employee_timelines = self.detect_conflicts(processed_events)
        
        # NUEVO: Detectar conexiones de viaje
        travel_connections = self.detect_travel_connections(processed_events)
        
        # Añadir info de viajes a cada evento
        for event in processed_events:
            event_id = event['event_id']
            if event_id in travel_connections:
                event['people_with_travel'] = travel_connections[event_id]['people_with_travel']
                event['travel_from_previous'] = travel_connections[event_id]['from_previous']
                event['travel_to_next'] = travel_connections[event_id]['to_next']
                
                # Marcar reservations con flag de viaje
                for res in event['reservations']:
                    res['has_travel_connection'] = res['employee'] in travel_connections[event_id]['people_with_travel']
        
        result = {
            'events': processed_events,
            'unassigned_events': unassigned_events,
            'stats': stats,
            'conflicts': conflicts,
            'employee_timelines': dict(employee_timelines),
            'last_updated': datetime.now().strftime('%d/%m/%Y %H:%M'),
            'now_date': datetime.now().date()
        }
        
        logger.info(f"✅ Procesados {stats['total_events']} eventos")
        return result
    
    def _determine_set(self, championship: str) -> str:
        """Determinar SET por campeonato"""
        if not championship:
            return 'default'
        
        championship_upper = championship.upper()
        
        for key, value in self.championship_to_set.items():
            if key in championship_upper:
                return value
        
        return 'default'
    
    def create_sharepoint_excel(self, processed_data: Dict) -> bool:
        """Crear Excel en SharePoint"""
        token = self.get_graph_token()
        if not token:
            logger.warning("⚠️ No se pudo obtener token de Graph, omitiendo creación de Excel")
            return False
        
        # Implementación de creación Excel...
        # (Por implementar completamente si lo necesitas)
        
        logger.info("✅ Excel creado en SharePoint (placeholder)")
        return True


# Aplicación Flask
app = Flask(__name__)
app.secret_key = os.environ.get('SECRET_KEY', 'aks-calendar-2025')

# Variables globales
calendar_instance = None
cached_dashboard_data = None
last_update_status = {'success': False, 'timestamp': None}

@app.route('/')
def dashboard():
    """Dashboard visual principal"""
    global calendar_instance, cached_dashboard_data
    
    if not calendar_instance:
        return render_template('config_needed.html')
    
    if not cached_dashboard_data:
        cached_dashboard_data = calendar_instance.process_motorsport_data()
    
    data = cached_dashboard_data
    
    if not data:
        return "<h1>Error obteniendo datos</h1>", 500
    
    # Renderizar template con datos
    return render_template('dashboard.html',
        stats=data['stats'],
        events=data['events'],
        unassigned_events=data['unassigned_events'],
        conflicts=data['conflicts'],
        last_updated=data['last_updated'],
        now_date=data['now_date'],
        color_mapping=calendar_instance.color_mapping
    )

@app.route('/config', methods=['GET', 'POST'])
def config():
    """Configuración del sistema"""
    if request.method == 'POST':
        try:
            config_data = {
                'airtable_token': request.form.get('airtable_token'),
                'airtable_base_id': request.form.get('airtable_base_id', 'app4p2TY96NofXW4u'),
                'tenant_id': request.form.get('tenant_id'),
                'client_id': request.form.get('client_id'),
                'client_secret': request.form.get('client_secret'),
                'sharepoint_site_url': request.form.get('sharepoint_site_url'),
                'auto_update_interval': 15
            }
            
            global calendar_instance
            calendar_instance = EventsCalendarAKS(config_data)
            
            return jsonify({'success': True, 'message': 'Sistema configurado'})
            
        except Exception as e:
            return jsonify({'error': str(e)}), 500
    
    return render_template('config.html')

@app.route('/update')
def manual_update():
    """Actualización manual"""
    global calendar_instance, cached_dashboard_data
    
    if not calendar_instance:
        return "Sistema no configurado", 400
    
    try:
        # Limpiar cache
        calendar_instance.cache = {}
        calendar_instance.cache_expiry = {}
        
        # Reprocesar datos
        cached_dashboard_data = calendar_instance.process_motorsport_data()
        
        if cached_dashboard_data:
            return """
            <html><head><meta charset="UTF-8"><meta http-equiv="refresh" content="2;url=/"></head>
            <body style="font-family: sans-serif; text-align: center; padding: 50px;">
                <h1>✅ Actualización exitosa</h1>
                <p>Redirigiendo...</p>
            </body>
            </html>
            """
        else:
            return "<h1>Error en actualización</h1>", 500
    except Exception as e:
        return f"<h1>Error: {str(e)}</h1>", 500

@app.route('/api/available-staff')
def api_available_staff():
    """API para buscar personal disponible"""
    global calendar_instance
    
    if not calendar_instance:
        return jsonify({'error': 'Sistema no configurado'}), 400
    
    start_date_str = request.args.get('start_date')
    end_date_str = request.args.get('end_date')
    role = request.args.get('role')
    
    if not start_date_str or not end_date_str:
        return jsonify({'error': 'Faltan fechas'}), 400
    
    try:
        start_date = datetime.strptime(start_date_str, '%Y-%m-%d').date()
        end_date = datetime.strptime(end_date_str, '%Y-%m-%d').date()
    except:
        return jsonify({'error': 'Formato de fecha inválido'}), 400
    
    available = calendar_instance.find_available_staff(start_date, end_date, role)
    
    return jsonify({
        'success': True,
        'count': len(available),
        'date_range': {
            'start': start_date_str,
            'end': end_date_str
        },
        'staff': available
    })

@app.route('/api/event-details/<event_id>')
def api_event_details(event_id):
    """API para obtener detalles completos de un evento"""
    global calendar_instance, cached_dashboard_data
    
    if not calendar_instance or not cached_dashboard_data:
        return jsonify({'error': 'Sistema no configurado'}), 400
    
    try:
        # Buscar el evento
        target_event = None
        for event in cached_dashboard_data['events']:
            if event['event_id'] == event_id:
                target_event = event
                break
        
        if not target_event:
            return jsonify({'error': 'Evento no encontrado'}), 404
        
        # 1. Info básica del evento
        event_info = {
            'event_id': target_event['event_id'],
            'event_name': target_event['event_name'],
            'city': target_event['city'],
            'set_name': target_event['set_name'],
            'color': target_event['color'],
            'coordinator': target_event['coordinator'],
            'from_date': target_event['from_date'].strftime('%d/%m/%Y'),
            'to_date': target_event['to_date'].strftime('%d/%m/%Y'),
            'duration_days': target_event['duration_days']
        }
        
        # 2. Personal asignado
        staff = []
        for res in target_event['reservations']:
            # Verificar si tiene conflictos
            has_conflict = False
            for conflict in cached_dashboard_data['conflicts']:
                if conflict['employee'] == res['employee']:
                    has_conflict = True
                    break
            
            staff.append({
                'name': res['employee'],
                'from_date': res['from_date'].strftime('%d/%m/%Y'),
                'to_date': res['to_date'].strftime('%d/%m/%Y'),
                'remote': res['remote'],
                'has_conflict': has_conflict
            })
        
        # 3. Eventos simultáneos (mismas fechas)
        simultaneous_events = []
        for event in cached_dashboard_data['events']:
            if event['event_id'] == event_id:
                continue
            
            # Verificar si hay solapamiento de fechas
            if not (event['to_date'] < target_event['from_date'] or 
                    event['from_date'] > target_event['to_date']):
                
                # Buscar personal compartido
                shared_staff = []
                for res in event['reservations']:
                    for target_res in target_event['reservations']:
                        if res['employee'] == target_res['employee']:
                            shared_staff.append(res['employee'])
                
                simultaneous_events.append({
                    'event_id': event['event_id'],
                    'event_name': event['event_name'],
                    'city': event['city'],
                    'set_name': event['set_name'],
                    'color': event['color'],
                    'from_date': event['from_date'].strftime('%d/%m/%Y'),
                    'to_date': event['to_date'].strftime('%d/%m/%Y'),
                    'shared_staff': shared_staff
                })
        
        # 4. Evento anterior más cercano
        previous_event = None
        min_days_before = float('inf')
        for event in cached_dashboard_data['events']:
            if event['to_date'] < target_event['from_date']:
                days_diff = (target_event['from_date'] - event['to_date']).days
                if days_diff < min_days_before:
                    min_days_before = days_diff
                    previous_event = {
                        'event_id': event['event_id'],
                        'event_name': event['event_name'],
                        'city': event['city'],
                        'set_name': event['set_name'],
                        'color': event['color'],
                        'from_date': event['from_date'].strftime('%d/%m/%Y'),
                        'to_date': event['to_date'].strftime('%d/%m/%Y'),
                        'days_before': days_diff
                    }
        
        # 5. Evento siguiente más cercano
        next_event = None
        min_days_after = float('inf')
        for event in cached_dashboard_data['events']:
            if event['from_date'] > target_event['to_date']:
                days_diff = (event['from_date'] - target_event['to_date']).days
                if days_diff < min_days_after:
                    min_days_after = days_diff
                    next_event = {
                        'event_id': event['event_id'],
                        'event_name': event['event_name'],
                        'city': event['city'],
                        'set_name': event['set_name'],
                        'color': event['color'],
                        'from_date': event['from_date'].strftime('%d/%m/%Y'),
                        'to_date': event['to_date'].strftime('%d/%m/%Y'),
                        'days_after': days_diff
                    }
        
        # 6. Análisis de viajes
        travel_analysis = {
            'has_previous': previous_event is not None,
            'has_next': next_event is not None,
            'days_from_previous': min_days_before if previous_event else None,
            'days_to_next': min_days_after if next_event else None
        }
        
        return jsonify({
            'success': True,
            'event': event_info,
            'staff': staff,
            'simultaneous_events': simultaneous_events,
            'previous_event': previous_event,
            'next_event': next_event,
            'travel_analysis': travel_analysis
        })
        
    except Exception as e:
        logger.error(f"Error obteniendo detalles de evento: {str(e)}")
        return jsonify({'error': str(e)}), 500

if __name__ == "__main__":
    logger.info("🏁 Events Calendar AKS - Al Kamel Management")
    
    port = int(os.environ.get('PORT', 5000))
    
    app.run(
        host='0.0.0.0',
        port=port,
        debug=False,
        threaded=True
    )